"""Importa il database completo (Abilità, Incantesimi, Miracoli, Rune, Cerchi, tabelle speciali) dal foglio Excel dello staff.

Uso: python3 -I scripts/importa_catalogo.py <file.xlsx> <cartella src/data>
Scrive catalogo.json (abilità), spells.json (tutti i poteri) e tabelle.json (fogli speciali).
Gli id già usati nelle schede restano invariati.
"""
import json, re, sys, unicodedata
import openpyxl

XLSX, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(XLSX, data_only=True)
SALTA = {'IN LAVORAZIONE - CdS', 'Foglio17'}


def t(x):
    """testo pulito; '\\' o '/' da soli valgono 'vuoto'"""
    if x is None:
        return ''
    s = str(x).replace('\r', '').strip()
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r' *\n *', '\n', s)
    return '' if s in ('\\', '/') else s


def slug(s):
    s = unicodedata.normalize('NFD', s.lower().replace('’', '').replace("'", ''))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')


def norm(s):
    return slug(s).replace('-', ' ')


# ------------------------------------------------------------------ abilità
LIV = ['Minore', 'Maggiore', 'Estremo', 'Assoluto', 'Puro']
LIV_F = ['Minore', 'Maggiore', 'Estrema', 'Assoluta', 'Pura']  # Conoscenza Alchemica
CAT = ['fisico', 'mentale', 'sovrannaturale']

# id storici (abilities.ts) -> nome nel foglio
STORICI = {
    'armi-corte': 'Armi Corte', 'armi-medie': 'Armi Medie', 'armi-due-mani': 'Armi a Due Mani', 'armi-asta': 'Armi in Asta',
    'armi-lancio': 'Armi da Lancio', 'armi-tiro': 'Armi da Tiro', 'scudo': 'Scudo', 'armatura-cuoio': 'Armatura di Cuoio',
    'armatura-cuoio-lavorato': 'Armatura di Cuoio Lavorato', 'armatura-anelli': 'Armatura ad Anelli',
    'armatura-piastre': 'Armatura di Piastre', 'armatura-completa': 'Armatura Completa', 'sviluppo-fisico-1': 'Sviluppo Fisico 1',
    'contare': 'Contare', 'leggere-scrivere': 'Leggere & Scrivere', 'stabilizzare': 'Stabilizzare',
    'conoscenza-erboristica': 'Conoscenza Erboristica', 'conoscenza-metallurgica': 'Conoscenza Metallurgica',
    'valutare-1': 'Valutare 1', 'ambidestria': 'Ambidestria', 'linguaggio-arcano': 'Linguaggio Arcano',
    'linguaggio-runico': 'Linguaggio Runico', 'percezione-sovrannaturale': 'Percezione Sovrannaturale',
    'conoscenze-magiche': 'Conoscenze Magiche', 'dono-arcano-1': 'Dono Arcano Minore', 'legame-divino-1': 'Legame Divino Minore',
    'talento-1': 'Talento 1', 'aett-3': '3^ Aett', 'aett-2': '2^ Aett', 'aett-1': '1^ Aett', 'efaett': '2^ Efaett',
}
for i, l in enumerate(LIV):
    STORICI[f'studio-arcano-{i + 1}'] = f'Studio Arcano {l}'
    STORICI[f'alchimia-{i + 1}'] = f'Conoscenza Alchemica {LIV_F[i]}'
    if i:
        STORICI[f'legame-divino-{i + 1}'] = f'Legame Divino {l}'
        if i < 4: STORICI[f'dono-arcano-{i + 1}'] = f'Dono Arcano {l}'


def potere(nome):
    """via e livello delle abilità di potere (per PP/TdR e per la scelta degli incantesimi)"""
    m = re.match(r'^(Studio Arcano|Dono Arcano|Legame Divino|Conoscenza Alchemica) (\w+)', nome)
    if m and '+1' not in nome:
        lv = m.group(2)
        n = (LIV.index(lv) if lv in LIV else LIV_F.index(lv) if lv in LIV_F else -1) + 1
        if n:
            via = {'Studio Arcano': 'mago', 'Dono Arcano': 'stregone', 'Legame Divino': 'sacerdote', 'Conoscenza Alchemica': 'alchimista'}[m.group(1)]
            return {'via': via, 'livello': n}
    if nome == 'Conoscenze Magiche':
        return {'via': 'mago', 'livello': 1}
    m = re.match(r'^([123])\^ Aett$', nome)
    if m:
        return {'via': 'runista', 'livello': 4 - int(m.group(1))}
    if re.match(r'^[12]\^ Efaett$', nome):
        return {'via': 'runista', 'livello': 4}
    return None


def tempo(x):
    """'\\' e 'x' = nessun Tempo di Apprendimento; 'NI' = non insegnabile; '/' o vuoto = disciplina non usata"""
    s = '' if x is None else str(x).strip()
    if s in ('', '/'):
        return None
    if s.lower() in ('x', '\\'):
        return 0
    if s.upper() == 'NI':
        return 'NI'
    try:
        return int(float(s))
    except ValueError:
        return None


righe = [r for r in wb['Abilità'].iter_rows(min_row=2, values_only=True) if t(r[1])]
abilita = []
for r in righe:
    nome = t(r[1])
    spec = t(r[7]).upper().replace('+', ' ')
    pc = re.match(r'^(\d+)\s*PC$', spec)
    lettere = set(re.findall(r'[SRTL]', spec)) if not pc else set()
    tempi = dict(zip(CAT, (tempo(r[4]), tempo(r[5]), tempo(r[6]))))
    usati = [c for c in CAT if tempi[c] is not None]
    numeri = [v for v in tempi.values() if isinstance(v, int)]
    gruppo = re.sub(r'\s*\n\s*', ', ', t(r[9]))
    note = '\n'.join(x for x in (t(r[11]), t(r[12]), t(r[13])) if x and x.lower() not in ('no', '?'))
    a = {
        'nome': nome, 'codice': t(r[10]), 'categoria': usati[0] if usati else 'mentale',
        'tipo': 'base' if gruppo == 'Base' else 'avanzata',
        'pc': int(pc.group(1)) if pc else None,
        'tda': max(numeri) if numeri else None,
        'tempi': {k: v for k, v in tempi.items() if v is not None},
        'sephira': int(t(r[0])) if t(r[0]).isdigit() else None,
        'titolo': t(r[2]), 'prerequisiti': t(r[3]), 'descrizione': t(r[8]), 'gruppo': gruppo, 'note': note,
    }
    if 'S' in lettere: a['segreta'] = True
    if 'R' in lettere: a['riacquistabile'] = True
    if 'T' in lettere: a['temporanea'] = True
    if 'L' in lettere: a['limitata'] = True
    if 'NI' in tempi.values(): a['nonInsegnabile'] = True
    if a['tipo'] == 'base' and a.get('nonInsegnabile'): a['soloCreazione'] = True
    p = potere(nome)
    if p: a['effetti'] = {'potere': p}
    if re.match(r'^Sviluppo Fisico \d$', nome): a['effetti'] = {'pf': 1}
    abilita.append({k: v for k, v in a.items() if v not in ('', None, {})})


def preferita(cands):
    def punti(a):
        return (a.get('gruppo') == 'Base', a.get('titolo') != 'Professione Segreta', not a.get('gruppo', '').startswith('Avanzamento'))
    return max(cands, key=punti)


per_nome = {}
for a in abilita:
    per_nome.setdefault(norm(a['nome']), []).append(a)
for id_, nome in STORICI.items():
    c = per_nome.get(norm(nome))
    if c:
        preferita(c)['id'] = id_
    else:
        print('ATTENZIONE: id storico senza riga nel foglio:', id_, nome, file=sys.stderr)
for a in abilita:
    if 'id' in a:
        continue
    base = slug(a['nome'].replace('+', ' piu '))
    p = a.get('effetti', {}).get('potere')
    # varianti delle abilità di potere: stesso prefisso dell'id storico (serve agli avvisi email)
    if p and p['via'] == 'sacerdote' and a['nome'].startswith('Legame Divino'): base = f"legame-divino-{p['livello']}"
    if p and p['via'] == 'mago' and a['nome'].startswith('Studio Arcano'): base = f"studio-arcano-{p['livello']}"
    a['id'] = f"{base}-{slug(a.get('codice') or 'senza-codice')}"
ids = [a['id'] for a in abilita]
assert len(ids) == len(set(ids)), 'id abilità duplicati'
abilita = [{'id': a['id'], **{k: v for k, v in a.items() if k != 'id'}} for a in abilita]

# ------------------------------------------------------------------ poteri
SCUOLA = {'Madre della Vita': 'Madre', "Padre dell'Oblio": 'Padre', 'Padre dell’Oblio': 'Padre', 'Spada della Giustizia': 'Spada',
          'Generici': 'Generico', 'F. Centrale': 'Centrale'}


def livello(s):
    s = t(s)
    m = re.match(r"^(\d\^ Aett)", s)
    if m:
        return m.group(1)
    return re.sub(r'\s*\(.*$', '', s).strip()


def segreto(s):
    s = t(s).strip().lower()
    if s.startswith('brevetto'):
        return {'segreto': True, 'brevetto': t(s).strip().capitalize()}
    return {'segreto': True} if s in ('si', 'sì') else {}


poteri = []
def add_potere(tipo, row, nome, scuola, liv, flavor, formula, aggiunta, descr, gitt, bers, lancio, durata, seg, codice, nota, extra=None):
    p = {'nome': nome, 'tipo': tipo, 'scuola': SCUOLA.get(scuola, scuola), 'livello': liv, 'flavor': flavor, 'formula': formula,
         'aggiunta': aggiunta, 'descrizione': descr, 'gittata': gitt, 'bersaglio': bers, 'lancio': lancio, 'durata': durata, 'codice': codice}
    p.update(extra or {})
    p.update(segreto(seg))
    if liv == 'Master': p['master'] = True
    if nota: p['notaStaff'] = nota
    poteri.append({k: v for k, v in p.items() if v not in (None,)})


for sh, tipo in (('Incantesimi', 'incantesimo'), ('Miracoli', 'miracolo')):
    for r in wb[sh].iter_rows(min_row=2, values_only=True):
        if not t(r[0]): continue
        add_potere(tipo, r, t(r[0]), t(r[2]), livello(r[3]), t(r[6]), t(r[7]), t(r[8]), t(r[9]), t(r[10]), t(r[11]), t(r[12]), t(r[13]),
                   r[14], t(r[15]), '; '.join(t(x) for x in r[16:19] if t(x)), {'genere': t(r[4]), 'breve': t(r[5])})
for r in wb['Rune'].iter_rows(min_row=2, values_only=True):
    if not t(r[1]): continue
    add_potere('runa', r, t(r[1]), t(r[2]), livello(r[3]), t(r[6]), t(r[7]), t(r[8]), t(r[9]), t(r[10]), t(r[11]), t(r[12]), t(r[13]),
               r[14], t(r[16]), '', {'genere': t(r[4]), 'lettera': t(r[0]), 'significato': t(r[5])})
for r in wb['Cerchi Alchemici'].iter_rows(min_row=2, values_only=True):
    if not t(r[0]): continue
    add_potere('cerchio', r, t(r[0]), t(r[2]), livello(r[3]), '', t(r[6]), t(r[7]), t(r[8]), t(r[9]), t(r[10]), t(r[11]), t(r[12]),
               r[13], t(r[14]), '', {'genere': t(r[4]), 'breve': t(r[5]), **({'proibito': True} if t(r[1]) == 'Proibito' else {})})

for p in poteri:
    p['id'] = re.sub(r'[^a-z0-9]+', '-', (p['tipo'][:3] + ' ' + p['scuola'] + ' ' + p['nome']).lower().replace('’', '').replace("'", '')).strip('-')
visti = {}
for p in poteri:
    if p['id'] in visti:
        p['id'] += '-' + slug(p['codice'] or str(len(visti)))
    visti[p['id']] = p
assert len(visti) == len(poteri), 'id poteri duplicati'
poteri = [{'id': p['id'], **{k: v for k, v in p.items() if k != 'id'}} for p in poteri]

# ------------------------------------------------------------------ tabelle speciali
FOGLI_SPECIALI = ['Linfa Eterea', 'Risonanza', 'Emozioni Spiriti', 'Messale', 'Rosario Inquisitorio', 'Ambasciata', 'Eco',
                  'Comunione con la Bestia', 'Forza del Branco', 'Robustezza Arcana', 'Robustezza Arcana Migliorata', 'Canalizzare Elemento']
tabelle = []
for nome in FOGLI_SPECIALI:
    grid = [[t(c) for c in r] for r in wb[nome].iter_rows(values_only=True)]
    grid = [r for r in grid if any(r)]
    w = max(max((i + 1 for i, c in enumerate(r) if c), default=0) for r in grid)
    grid = [r[:w] for r in grid]
    while grid and all(not r[0] for r in grid):
        grid = [r[1:] for r in grid]
    # prima riga con il solo titolo del foglio: il titolo lo mostra già l'app
    if grid and sum(1 for c in grid[0] if c) == 1 and norm(next(c for c in grid[0] if c)) == norm(nome):
        grid = grid[1:]
    tabelle.append({'nome': nome, 'righe': grid})
# abilità collegate: il nome della tabella è l'inizio del nome dell'abilità (vince il più lungo)
for a in abilita:
    n = norm(a['nome'])
    cand = [tb['nome'] for tb in tabelle if n == norm(tb['nome']) or n.startswith(norm(tb['nome']) + ' ')]
    if cand:
        best = max(cand, key=len)
        a['tabelle'] = [best] + (['Emozioni Spiriti'] if best in ('Linfa Eterea', 'Risonanza') else [])

json.dump(abilita, open(f'{OUT}/catalogo.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
json.dump(poteri, open(f'{OUT}/spells.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
json.dump(tabelle, open(f'{OUT}/tabelle.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
print(len(abilita), 'abilità,', len(poteri), 'poteri,', len(tabelle), 'tabelle')
print('collegate a tabelle:', sorted({(a['nome'], tuple(a['tabelle'])) for a in abilita if 'tabelle' in a}))
